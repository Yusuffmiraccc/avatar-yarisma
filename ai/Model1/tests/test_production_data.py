import csv
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

import pytest

from app.production_data import (
    COST_COMPONENT_FIELDS,
    CSV_COLUMNS,
    MODEL1_FEATURE_MAP,
    calculate_cost_total,
    calculate_ekp,
    model1_scenario_a_features,
    validate_records,
)
from training.validate_ekp_chronological import chronological_splits, evaluate_records


def synthetic_test_fixture(**changes):
    """Explicitly synthetic in-memory test fixture; never saved as factory data."""
    row = {
        "batch_id": "SENTETIK-ORNEK-TEST-001",
        "source_type": "synthetic",
        "source_system": "SENTETIK ORNEK — pytest fixture",
        "source_reference": "SENTETIK ORNEK — not a factory record",
        "validation_status": "pending",
        "production_started_at": "2026-01-01T08:00:00+03:00",
        "measurement_at": "2026-01-01T18:00:00+03:00",
        "measurement_method": "SENTETIK ORNEK scale fixture",
        "currency": "TRY",
        "cost_period": "per_batch",
        "total_cost_tl": "1200.00",
        "net_output_kg": "10.00",
    }
    row.update(changes)
    return row


def issue_codes(report):
    return {issue["code"] for issue in report["issues"]}


def test_cost_total_rejects_duplicate_and_negative_components():
    assert calculate_cost_total([
        {"name": "meat", "amount_tl": "100.20"},
        {"name": "energy", "amount_tl": "8.30"},
    ]) == Decimal("108.50")
    with pytest.raises(ValueError, match="Duplicate"):
        calculate_cost_total([
            {"name": "meat", "amount_tl": 1},
            {"name": "Meat", "amount_tl": 2},
        ])
    with pytest.raises(ValueError, match="Negative"):
        calculate_cost_total([{"name": "meat", "amount_tl": -1}])


def test_ekp_formula_units_reconciliation_and_zero_output():
    assert calculate_ekp(Decimal("1200.00"), Decimal("10.00")) == Decimal("120")
    parts = [{"name": "meat", "amount_tl": 100}, {"name": "fason", "amount_tl": 20}]
    assert calculate_ekp(None, 10, components=parts) == Decimal("12")
    assert calculate_ekp(120, 10, components=parts) == Decimal("12")
    with pytest.raises(ValueError, match="greater than zero"):
        calculate_ekp(100, 0)
    with pytest.raises(ValueError, match="TRY"):
        calculate_ekp(100, 1, currency="USD")
    with pytest.raises(ValueError, match="reconcile"):
        calculate_ekp(200, 10, components=parts)


def test_valid_record_is_kept_and_ekp_is_calculated_but_not_time_evaluation_eligible():
    report = validate_records([synthetic_test_fixture()])
    assert report["rows"] == 1
    assert report["error_count"] == 0
    assert report["records"][0]["calculated_ekp_tl_kg"] == "120"
    assert report["records"][0]["eligible_for_ekp"] is True
    assert report["records"][0]["eligible_for_time_evaluation"] is False


def test_required_unknown_and_wrong_type_fields_are_reported_by_row_and_field():
    row = synthetic_test_fixture()
    row.pop("measurement_method")
    row["source_system"] = 99
    row["unexpected"] = "value"
    report = validate_records([row])
    problems = {(issue["field"], issue["code"]) for issue in report["issues"]}
    assert ("measurement_method", "required") in problems
    assert ("source_system", "wrong_type") in problems
    assert ("unexpected", "unknown_field") in problems
    assert report["records"][0]["is_valid"] is False


@pytest.mark.parametrize("field,value,code", [
    ("total_cost_tl", "NaN", "invalid_number"),
    ("net_output_kg", "Infinity", "invalid_number"),
    ("meat_input_kg", "-0.1", "negative"),
    ("drying_loss_pct", "101", "out_of_range"),
    ("net_output_kg", "0", "nonpositive"),
])
def test_nonfinite_negative_percentage_and_zero_values_are_rejected(field, value, code):
    report = validate_records([synthetic_test_fixture(**{field: value})])
    assert code in issue_codes(report)


def test_duplicate_batch_ids_are_rejected_without_dropping_rows():
    report = validate_records([
        synthetic_test_fixture(),
        synthetic_test_fixture(batch_id="SENTETIK-ORNEK-TEST-001"),
    ])
    assert report["rows"] == 2
    assert sum(i["code"] == "duplicate_batch_id" for i in report["issues"]) == 2


def test_timezone_and_chronology_errors_are_reported():
    naive = synthetic_test_fixture(production_started_at="2026-01-01T08:00:00")
    assert "invalid_timestamp" in issue_codes(validate_records([naive]))
    aware = synthetic_test_fixture(
        production_started_at="2026-01-01T08:00:00+03:00",
        measurement_at="2026-01-01T07:00:00+03:00",
        production_ended_at="2026-01-01T06:00:00+03:00",
    )
    assert "time_order" in issue_codes(validate_records([aware]))


def test_complete_cost_breakdown_requires_every_component_and_reconciles():
    row = synthetic_test_fixture(cost_breakdown_complete="true", meat_cost_tl="1200.00")
    report = validate_records([row])
    assert "incomplete_cost_breakdown" in issue_codes(report)

    complete = synthetic_test_fixture(cost_breakdown_complete=True)
    for field in COST_COMPONENT_FIELDS:
        complete[field] = "0.00"
    complete["meat_cost_tl"] = "1200.00"
    report = validate_records([complete])
    assert "cost_reconciliation" not in issue_codes(report)
    complete["other_cost_tl"] = "2.00"
    assert "cost_reconciliation" in issue_codes(validate_records([complete]))


def test_spice_item_details_reconcile_without_double_counting_spice_total():
    detail = '[{"name":"kimyon","quantity_kg":2,"unit_price_tl_kg":10,"cost_tl":20,"source_reference":"SENTETIK ORNEK"}]'
    row = synthetic_test_fixture(
        spice_details_json=detail,
        spice_details_complete=True,
        spice_cost_tl="20.00",
    )
    assert issue_codes(validate_records([row])).isdisjoint({"invalid_json", "spice_item_reconciliation", "spice_detail_reconciliation"})
    row["spice_cost_tl"] = "21.00"
    assert "spice_detail_reconciliation" in issue_codes(validate_records([row]))
    row["spice_details_json"] = '[{"name":"kimyon","quantity_kg":NaN,"unit_price_tl_kg":10,"cost_tl":20}]'
    assert "invalid_json" in issue_codes(validate_records([row]))


def test_mass_balance_requires_measured_losses_and_reports_mismatch():
    row = synthetic_test_fixture(
        input_mass_kg="15.000",
        net_output_kg="10.000",
        drying_loss_kg="4.000",
        sorting_loss_kg="1.000",
        other_loss_kg="0.000",
        losses_complete=True,
        loss_measurement_at="2026-01-01T18:00:00+03:00",
        loss_measurement_method="SENTETIK ORNEK scale fixture",
    )
    assert "mass_balance_mismatch" not in issue_codes(validate_records([row]))
    row["drying_loss_kg"] = "3.000"
    assert "mass_balance_mismatch" in issue_codes(validate_records([row]))


def test_output_above_recorded_input_is_a_boundary_warning_unless_complete_balance_fails():
    row = synthetic_test_fixture(input_mass_kg="10.00", net_output_kg="11.00")
    report = validate_records([row])
    assert "input_mass_check" in issue_codes(report)
    assert report["records"][0]["is_valid"] is True


def test_loss_percentages_never_enter_scenario_a_and_unknown_time_is_flagged():
    report = validate_records([synthetic_test_fixture(drying_loss_pct=4.2)])
    assert "unknown_loss_basis" in issue_codes(report)
    model_fields = set(MODEL1_FEATURE_MAP)
    assert not model_fields.intersection({"Kuruma_kaybi_pct", "Ayiklama_fire_pct", "drying_loss_pct"})


def test_preproduction_feature_builder_requires_frozen_preproduction_snapshot():
    row = synthetic_test_fixture(
        production_started_at="2026-01-01T08:00:00+03:00",
        plan_snapshot_at="2026-01-01T07:30:00+03:00",
    )
    for field in MODEL1_FEATURE_MAP.values():
        if field.startswith("planned_") and field not in {
            "planned_region", "planned_recipe_profile", "planned_process_method"
        }:
            row[field] = "1.0"
    row["planned_meat_price_tl_kg"] = "20"
    row["planned_region"] = "İç Anadolu"
    row["planned_recipe_profile"] = "MEB_STANDART"
    row["planned_process_method"] = "Fason"
    features = model1_scenario_a_features(row)
    assert list(features) == list(MODEL1_FEATURE_MAP)
    assert features["Et_girdi_kg"] == 1.0
    with pytest.raises(ValueError, match="after production"):
        model1_scenario_a_features({**row, "plan_snapshot_at": "2026-01-01T09:00:00+03:00"})


def test_chronological_splits_keep_timestamp_groups_together_and_ordered():
    rows = []
    start = datetime(2026, 1, 1, tzinfo=timezone.utc)
    for day in range(12):
        for unit in range(2):
            stamp = (start + timedelta(days=day)).isoformat()
            rows.append({
                "batch_id": f"SENTETIK-ORNEK-SPLIT-{day}-{unit}",
                "production_started_at": stamp,
            })
    splits = chronological_splits(rows, min_train_batches=4, test_timestamps=2, n_splits=3)
    for train_indices, test_indices in splits:
        train = [rows[i] for i in train_indices]
        test = [rows[i] for i in test_indices]
        assert set(r["batch_id"] for r in train).isdisjoint(r["batch_id"] for r in test)
        assert max(r["production_started_at"] for r in train) < min(r["production_started_at"] for r in test)
        assert len(test) % 2 == 0


def test_time_evaluation_refuses_empty_or_synthetic_data_without_metrics():
    empty = evaluate_records([])
    assert empty["status"] == "no_real_data"
    assert empty["metrics"] == {}

    blocked = evaluate_records([synthetic_test_fixture()])
    assert blocked["status"] == "blocked_non_real_or_mixed_source"
    assert blocked["metrics"] == {}


def test_schema_has_unique_column_names_and_dictionary_entries():
    assert len(CSV_COLUMNS) == len(set(CSV_COLUMNS))
    assert len(CSV_COLUMNS) == len({entry[0] for entry in __import__("app.production_data", fromlist=["FIELD_DICTIONARY"]).FIELD_DICTIONARY})


def test_validated_rows_need_an_outcome_approval_time():
    missing = synthetic_test_fixture(validation_status="validated")
    assert "required_for_validation" in issue_codes(validate_records([missing]))
    ordered = synthetic_test_fixture(
        validation_status="validated",
        record_created_at="2026-01-01T19:00:00+03:00",
        validated_at="2026-01-01T18:30:00+03:00",
    )
    assert "time_order" in issue_codes(validate_records([ordered]))


def test_generated_templates_are_header_only_and_xlsx_has_dictionary():
    root = Path(__file__).resolve().parents[1]
    csv_path = root / "data" / "templates" / "real_production_batches_template.csv"
    xlsx_path = root / "data" / "templates" / "real_production_batches_template.xlsx"
    with csv_path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        headers = next(reader)
        assert next(reader, None) is None
    assert headers == CSV_COLUMNS
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    assert workbook["Production_Batches"].max_row == 1
    assert workbook["Production_Batches"].max_column == len(CSV_COLUMNS)
    assert workbook["Data_Dictionary"].max_row == len(CSV_COLUMNS) + 1
    assert "örnek üretim kaydı içermez" in " ".join(
        str(cell.value or "") for row in workbook["Instructions"].iter_rows() for cell in row
    ).casefold()
