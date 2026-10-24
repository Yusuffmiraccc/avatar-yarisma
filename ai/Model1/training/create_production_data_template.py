"""Generate blank CSV/XLSX templates and dictionary; never emits example production records."""
from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.production_data import CSV_COLUMNS, FIELD_DICTIONARY

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_DIR = ROOT / "data" / "templates"
CSV_PATH = TEMPLATE_DIR / "real_production_batches_template.csv"
XLSX_PATH = TEMPLATE_DIR / "real_production_batches_template.xlsx"
DICTIONARY_PATH = ROOT / "reports" / "production_data_dictionary.md"


def main() -> None:
    TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)
    with CSV_PATH.open("w", newline="", encoding="utf-8-sig") as file:
        csv.writer(file).writerow(CSV_COLUMNS)

    workbook = Workbook()
    instructions = workbook.active
    instructions.title = "Instructions"
    instruction_rows = [
        ["AVATAR Model1 — boş gerçek üretim verisi şablonu"],
        ["Bu çalışma kitabı boş şablondur; gerçek veya örnek üretim kaydı içermez."],
        ["Kaynak sınıfı gerçek olmayan kayıtları gerçek veri değerlendirmesiyle birleştirmeyin."],
        ["İsteğe bağlı/ölçülmemiş hücreleri boş bırakın; tahmin ederek doldurmayın."],
        ["Zaman damgaları ISO-8601 ve saat dilimi içermeli; biçim: YYYY-MM-DDTHH:MM:SS+03:00."],
        ["Miktarlar kg, metre, saat, kWh; maliyetler batch başına TRY; EKP TRY/kg olarak girilir."],
        ["Döviz dönüşümü yapılmaz. EKP yalnızca currency=TRY ve cost_period=per_batch için hesaplanır."],
        ["cost_breakdown_complete=true ise 10 ayrık maliyet kaleminin tümüne tutar girin; sıfır için 0 yazın."],
        ["spice_details_complete=true ise her baharat için ad, kg miktarı, TL/kg fiyatı ve TL maliyeti girin; bunlar spice_cost_tl içinde yalnızca bir kez sayılır."],
        ["losses_complete=true ise kayıp kg alanları, ölçüm zamanı ve yöntem girilmelidir."],
        ["Üretim öncesi model girdileri planned_* alanlarıdır; plan_snapshot_at üretim başlangıcından geç olamaz."],
        ["Gerçekleşmiş çıktı, maliyet ve kayıp oranları üretim öncesi model girdisi değildir."],
        ["validation_status=validated yalnızca kaynak sahibi kontrolünden sonra seçilmelidir."],
        ["Alan açıklamaları ve kaynak/timing bilgileri Data_Dictionary sayfasındadır."],
    ]
    for row in instruction_rows:
        instructions.append(row)
    instructions.column_dimensions["A"].width = 118
    instructions.freeze_panes = "A2"

    batches = workbook.create_sheet("Production_Batches")
    batches.append(CSV_COLUMNS)
    batches.freeze_panes = "A2"
    batches.auto_filter.ref = f"A1:{get_column_letter(len(CSV_COLUMNS))}1"
    for idx, name in enumerate(CSV_COLUMNS, start=1):
        batches.column_dimensions[get_column_letter(idx)].width = min(max(len(name) + 3, 15), 34)

    dictionary = workbook.create_sheet("Data_Dictionary")
    dictionary.append(["field", "description", "data_type", "unit", "required_for", "accepted_values", "source", "measurement_timing"])
    for name, data_type, unit, required, accepted, source_name, timing in FIELD_DICTIONARY:
        description = name.replace("_", " ") + " — " + required
        dictionary.append([name, description, data_type, unit, required, accepted, source_name, timing])
    dictionary.freeze_panes = "A2"
    dictionary.auto_filter.ref = f"A1:H{dictionary.max_row}"
    for column, width in {"A": 34, "B": 66, "C": 18, "D": 22, "E": 30, "F": 48, "G": 34, "H": 34}.items():
        dictionary.column_dimensions[column].width = width
    for sheet in workbook.worksheets:
        sheet.row_dimensions[1].height = 30
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="174A7E")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    workbook.save(XLSX_PATH)

    lines = [
        "# Gerçek üretim veri sözlüğü",
        "",
        "Bu sözlük boş CSV/XLSX şablonuyla aynı şemayı kullanır. Boş değer bilinmiyor demektir;",
        "eksik ölçümler tahmin edilerek doldurulmaz. CSV/XLSX örnek üretim kaydı içermez.",
        "",
        "| Alan | Açıklama | Tip | Birim | Zorunluluk/kullanım | Kabul edilen değerler | Kaynak | Ölçüm zamanı |",
        "|---|---|---|---|---|---|---|---|",
    ]
    for name, data_type, unit, required, accepted, source_name, timing in FIELD_DICTIONARY:
        description = name.replace("_", " ") + " — " + required
        escaped = [str(v).replace("|", "\\|").replace("\n", " ") for v in
                   (name, description, data_type, unit, required, accepted, source_name, timing)]
        lines.append("| " + " | ".join(escaped) + " |")
    lines.extend([
        "",
        "Maliyet alanları TRY/batch ve birbirinden ayrık kabul edilir. Aynı fatura kalemini iki alana",
        "yazmayın. Fason hizmeti ve kurum içi işleme ayrı maliyet alanlarıdır. Toplam maliyet, cost_breakdown_complete=true ise tüm on bileşenle",
        "kuruş yuvarlama toleransında mutabık olmalıdır. EKP = toplam parti maliyeti / ölçülmüş net",
        "çıktı (TRY/kg). Harici kur dönüşümü ve aylık maliyet tahsisi bu şemada yapılmaz.",
        "",
        "Üretim öncesi özellikler yalnız planned_* alanlarından oluşturulur ve plan_snapshot_at",
        "üretim başlangıcından geç olamaz. Gerçekleşmiş net çıktı, kayıp yüzdesi ve maliyet bileşenleri",
        "sonuç alanlarıdır; Scenario A girdilerinden hariç tutulur.",
    ])
    DICTIONARY_PATH.parent.mkdir(parents=True, exist_ok=True)
    DICTIONARY_PATH.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"Created blank template: {CSV_PATH}")
    print(f"Created blank workbook: {XLSX_PATH}")
    print(f"Created dictionary: {DICTIONARY_PATH}")


if __name__ == "__main__":
    main()
